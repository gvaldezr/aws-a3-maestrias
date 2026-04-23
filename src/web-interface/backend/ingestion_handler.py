"""
Document Ingestion Handler — U7: Parsea documentos y dispara el pipeline.
Lambda invocado por S3 Event (PutObject en uploads/).
"""
from __future__ import annotations

import io
import json
import os
import uuid
from datetime import datetime, timezone
from typing import Any

import boto3

from src.infrastructure.observability.logger import get_logger
from src.infrastructure.observability.metrics import record_metric
from src.infrastructure.state.models import StateMetadata, SubjectState
from src.infrastructure.state.state_manager import update_subject_state
from src.web_interface.backend.document_parser import parse_text_to_document

logger = get_logger(__name__)


def _extract_text_from_s3(bucket: str, key: str) -> str:
    """Descarga el archivo de S3 y extrae texto según el formato."""
    s3 = boto3.client("s3")
    obj = s3.get_object(Bucket=bucket, Key=key)
    file_bytes = obj["Body"].read()
    ext = key.rsplit(".", 1)[-1].lower()

    if ext == "pdf":
        return _extract_pdf(file_bytes)
    elif ext == "docx":
        return _extract_docx(file_bytes)
    elif ext in ("xlsx", "xls"):
        return _extract_xlsx(file_bytes)
    return file_bytes.decode("utf-8", errors="ignore")


def _extract_pdf(file_bytes: bytes) -> str:
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)
    except ImportError:
        return file_bytes.decode("utf-8", errors="ignore")


def _extract_docx(file_bytes: bytes) -> str:
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except ImportError:
        return file_bytes.decode("utf-8", errors="ignore")


def _extract_xlsx(file_bytes: bytes) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True)
        lines = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                line = " | ".join(str(c) for c in row if c is not None)
                if line.strip():
                    lines.append(line)
        return "\n".join(lines)
    except ImportError:
        return file_bytes.decode("utf-8", errors="ignore")


def _build_initial_json(parsed: "ParsedDocument") -> dict:
    """Construye el JSON inicial de la asignatura con estado INGESTED."""
    now = datetime.now(timezone.utc).isoformat()
    return {
        "subject_id": parsed.subject_id,
        "schema_version": "1.0",
        "created_at": now,
        "updated_at": now,
        "metadata": {
            "subject_name": parsed.subject_name,
            "program_name": parsed.program_name,
            "program_type": parsed.program_type,
            "subject_type": parsed.subject_type,
            "language": parsed.language,
        },
        "academic_inputs": {
            "graduation_profile": parsed.graduation_profile,
            "competencies": parsed.competencies,
            "learning_outcomes": parsed.learning_outcomes,
            "syllabus": parsed.syllabus,
        },
        "pipeline_state": {
            "current_state": "INGESTED",
            "state_history": [],
        },
    }


def _invoke_scholar_agent(subject_id: str) -> None:
    """Invoca el Agente Scholar en AgentCore Runtime para iniciar el pipeline."""
    runtime_arn = os.environ.get("SCHOLAR_AGENT_RUNTIME_ARN", "")
    if not runtime_arn:
        logger.warning("scholar_runtime_arn_not_configured", extra={"subject_id": subject_id})
        return

    bedrock_agentcore = boto3.client("bedrock-agentcore", region_name=os.environ.get("AWS_REGION", "us-east-1"))
    bedrock_agentcore.invoke_agent_runtime(
        agentRuntimeArn=runtime_arn,
        qualifier="DEFAULT",
        runtimeSessionId=subject_id,
        payload=json.dumps({"subject_id": subject_id}).encode(),
    )
    logger.info("scholar_agent_invoked", extra={"subject_id": subject_id})


def lambda_handler(event: dict, context: Any) -> dict:
    """
    Procesa eventos S3 PutObject: parsea el documento y dispara el pipeline.
    """
    records = event.get("Records", [])
    processed = 0

    for record in records:
        bucket = record["s3"]["bucket"]["name"]
        key = record["s3"]["object"]["key"]

        # Extraer subject_id del path: uploads/{subject_id}/{filename}
        parts = key.split("/")
        subject_id = parts[1] if len(parts) >= 3 else str(uuid.uuid4())

        logger.info("ingestion_start", extra={"subject_id": subject_id, "s3_key": key})

        try:
            text = _extract_text_from_s3(bucket, key)
            parsed = parse_text_to_document(text, subject_id, source_file=key)
            initial_json = _build_initial_json(parsed)

            # Persistir JSON inicial en S3 + DynamoDB
            from src.infrastructure.state.state_manager import save_subject_json
            save_subject_json(initial_json)

            # Disparar Agente Scholar (BR-W04)
            _invoke_scholar_agent(subject_id)

            record_metric("DocumentsIngested", 1, "Count")
            logger.info("ingestion_complete", extra={"subject_id": subject_id})
            processed += 1

        except Exception as exc:
            logger.error("ingestion_error", extra={"subject_id": subject_id, "error": str(exc)})
            record_metric("ParseErrors", 1, "Count")

    return {"statusCode": 200, "body": json.dumps({"processed": processed})}
